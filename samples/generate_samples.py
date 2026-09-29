#!/usr/bin/env python3
"""샘플 파일 생성기(개발용). 결과물은 저장소에 이미 들어 있으므로 사용자는 실행할 필요가 없다.
다시 만들고 싶을 때만: pip install openpyxl && python3 samples/generate_samples.py
모든 이름·생년월일·번호는 가상이다. 주민등록번호는 앞 7자리만 쓰고 나머지는 ***** 로 가렸다(시스템은 생년월일만 읽고 저장하지 않는다)."""
import struct
import zlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent

# ── 가상 훈련생(setup-sample-data.mjs·README 와 같은 값) ────────────────────────
TRAINEES = [  # (이름, 가상 생년월일, 가상 주민번호 앞부분)
    ('샘플가온', '1998-03-15', '980315-1******'),
    ('샘플나래', '1999-07-21', '990721-2******'),
    ('샘플다솜', '2000-01-09', '000109-4******'),
    ('샘플라온', '1997-11-30', '971130-1******'),
    ('샘플마루', '2001-05-05', '010505-3******'),
    ('샘플바다', '1996-09-12', '960912-2******'),
]
UNKNOWN = ('샘플없는사람', '900101-1******')  # 과정에 없는 사람(오류 행 시험용)

# ── 1) 공식 출석부(엑셀) ────────────────────────────────────────────────────────
# 열: 날짜 21~30일. 21·22일은 수업이 있는 날, 23·24일은 CSV 시험용이라 비워 둠, 25일은 회차가 없는 날(오류 시험),
# 26·27일 주말(빈칸), 28~30일 아직 오지 않은 날(-)
DAYS = list(range(21, 31))
WEEKDAY = {21: '월', 22: '화', 23: '수', 24: '목', 25: '금', 26: '토', 27: '일', 28: '월', 29: '화', 30: '수'}
MARKS = {  # 이름 → {일: 기호}
    '샘플가온': {21: '○', 22: '○', 25: '○'},
    '샘플나래': {21: '×', 22: '○', 25: '○'},
    '샘플다솜': {21: '◎', 22: '▶', 25: '○'},
    '샘플라온': {21: '×', 22: '▦15', 25: '○'},
    '샘플마루': {21: '▲', 22: '○', 25: '○'},
    '샘플바다': {21: '○', 22: '◎', 25: '○'},
    UNKNOWN[0]: {21: '○', 22: '○', 25: '○'},
}
for d in (28, 29, 30):
    for m in MARKS.values():
        m[d] = '-'

wb = Workbook()
ws = wb.active
ws.title = 'sheet1'  # 시스템은 첫 번째 시트만 읽는다
thin = Side(style='thin', color='999999')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
gray = PatternFill('solid', fgColor='E8E8E8')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

fixed = ['연번', '성명', '주민등록번호', '취약계층', '훈련생\n유형', '훈련생\n상태', '훈\n련\n일\n수', '출\n석\n일\n수', '결\n석\n일\n수', '휴\n가\n일\n수', '공\n가\n일\n수', '출석률\n일\n(%)', '출석률\n분\n(%)']
for i, title in enumerate(fixed, start=1):
    ws.cell(row=1, column=i, value=title)
    ws.merge_cells(start_row=1, start_column=i, end_row=3, end_column=i)
first_day_col = len(fixed) + 1  # N
ws.cell(row=1, column=first_day_col, value='2026년09월')
ws.merge_cells(start_row=1, start_column=first_day_col, end_row=1, end_column=first_day_col + len(DAYS) - 1)
last_col = first_day_col + len(DAYS)
ws.cell(row=1, column=last_col, value='비고')
ws.merge_cells(start_row=1, start_column=last_col, end_row=3, end_column=last_col)
for j, d in enumerate(DAYS):
    ws.cell(row=2, column=first_day_col + j, value=d)
    ws.cell(row=3, column=first_day_col + j, value=WEEKDAY[d])

people = [(n, r) for n, _b, r in TRAINEES] + [UNKNOWN]
for idx, (name, rrn) in enumerate(people, start=1):
    r = 3 + idx
    marks = MARKS[name]
    attended = sum(1 for v in marks.values() if v[0] in '○◎▲▶▦')
    absent = sum(1 for v in marks.values() if v == '×')
    excused = sum(1 for v in marks.values() if v.startswith('▦'))
    row = [idx, name, rrn, None, None, '훈련중', 19, attended, absent, 0, excused, None, None]
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v)
    for j, d in enumerate(DAYS):
        if d in marks:
            ws.cell(row=r, column=first_day_col + j, value=marks[d])
    ws.cell(row=r, column=last_col, value='가상 데이터(과정에 없는 사람)' if name == UNKNOWN[0] else None)
for row in ws.iter_rows(min_row=1, max_row=3 + len(people), min_col=1, max_col=last_col):
    for cell in row:
        cell.border = border
        cell.alignment = center
        if cell.row <= 3:
            cell.fill = gray
            cell.font = Font(bold=True)
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 18
ws.column_dimensions[ws.cell(row=1, column=last_col).column_letter].width = 26
ws.row_dimensions[1].height = 34

guide = wb.create_sheet('안내')
lines = [
    '가상 데이터 — 공식 출결 대사(S29) 시험용 출석부 샘플',
    '이 시트(안내)는 시스템이 읽지 않습니다. 시스템은 첫 번째 시트(sheet1)만 읽습니다.',
    '모든 이름·주민등록번호는 가상입니다(주민번호는 뒤 6자리를 가렸고, 시스템은 생년월일만 계산해 쓰며 저장하지 않습니다).',
    '기호: ○ 출석 / × 결석 / ◎ 지각 / ▲ 조퇴 / ▶ 외출(출석 처리) / ▦15 공가(인정결석) / - 아직 오지 않은 날 / 빈칸 수업 없는 날',
    '21·22일: 수업이 있는 날. 23·24일: CSV 시험용이라 비워 둠. 25일: 수업 회차가 없는 날(오류 시험). 26·27일: 주말.',
    '샘플없는사람: 과정에 등록되지 않은 사람(오류 행 시험).',
    '자세한 예상 결과는 samples/README.md 를 보세요.',
]
for i, t in enumerate(lines, start=1):
    guide.cell(row=i, column=1, value=t)
guide.column_dimensions['A'].width = 120
wb.save(ROOT / 'official-attendance' / 'official-attendance_sheet_sample.xlsx')

# ── 2) 공식 출결 CSV ────────────────────────────────────────────────────────────
csv_rows = [
    'round_no,trainee_name,birth_date,status,check_in,check_out',
    '3,샘플가온,1998-03-15,출석,09:10,17:55',      # 내부 09:00~18:00 과 15분 이내 → 공식으로 전환
    '3,샘플나래,1999-07-21,출석,09:40,18:00',      # 내부 09:00 과 40분 차이 → 확인 필요(시각 불일치)
    '3,샘플다솜,,지각,09:25,18:00',                # 내부 기록 없음 → 새로 기록
    '4,샘플라온,1997-11-30,결석,,',
    '4,샘플마루,,인정결석,,',
    '4,샘플바다,,PRESENT,09:00,18:00',             # 영문 상태 코드도 가능
    '4,샘플가온,1990-01-01,출석,,',                # 오류: 생년월일이 다른 사람
    '9,샘플나래,,출석,,',                          # 오류: 9회차 없음
    '4,샘플없는사람,,출석,,',                      # 오류: 과정에 없는 훈련생
    '4,샘플다솜,,모름,,',                          # 오류: 알 수 없는 상태
    '4,샘플라온,,출석,25:99,',                     # 오류: 시각 형식
    '4,샘플바다,,결석,,',                          # 오류: 같은 훈련생·회차 중복
]
(ROOT / 'official-attendance' / 'official-attendance_csv_sample.csv').write_bytes(('﻿' + '\r\n'.join(csv_rows) + '\r\n').encode('utf-8'))

# ── 3) 첨부 시험용 가상 파일 ───────────────────────────────────────────────────
def make_pdf() -> bytes:
    text = 'Sample deliverable (virtual data) - for upload test only'
    objs = [
        '<< /Type /Catalog /Pages 2 0 R >>',
        '<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 420 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
        None,
        '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    stream = f'BT /F1 14 Tf 20 100 Td ({text}) Tj ET'
    objs[3] = f'<< /Length {len(stream)} >>\nstream\n{stream}\nendstream'
    out = b'%PDF-1.4\n'
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f'{i} 0 obj\n{body}\nendobj\n'.encode('latin-1')
    xref = len(out)
    out += f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode()
    for off in offsets:
        out += f'{off:010d} 00000 n \n'.encode()
    out += f'trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    return out


def make_png(w: int = 240, h: int = 120) -> bytes:
    rows = bytearray()
    for y in range(h):
        rows.append(0)
        for x in range(w):
            border = x < 4 or y < 4 or x >= w - 4 or y >= h - 4
            stripe = ((x + y) // 12) % 2 == 0
            rows += bytes((60, 120, 200) if border else ((235, 240, 250) if stripe else (215, 225, 245)))

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack('>I', len(data)) + body + struct.pack('>I', zlib.crc32(body) & 0xFFFFFFFF)

    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(bytes(rows))) + chunk(b'IEND', b'')


(ROOT / 'attachments' / 'submission_sample.pdf').write_bytes(make_pdf())
(ROOT / 'attachments' / 'operation-log_photo_sample.png').write_bytes(make_png())
(ROOT / 'attachments' / 'operation-log_sample.txt').write_text(
    '가상 데이터 — 회차별 운영일지 첨부 시험용 파일\n'
    '교육장 프로젝터 점검 결과(가상): 이상 없음.\n'
    '이 파일은 업로드 시험용이며 실제 기관·개인과 무관합니다.\n',
    encoding='utf-8',
)
print('생성 완료:', *sorted(str(p.relative_to(ROOT)) for p in ROOT.rglob('*') if p.is_file() and p.name != 'generate_samples.py'), sep='\n  ')
